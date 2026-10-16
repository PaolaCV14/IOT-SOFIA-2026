#!/usr/bin/env python3
"""
==============================================================================
CONVERSOR DE MARKDOWN A WORD (.DOCX) ELEGANTE - PROYECTO SOFÍA 2026
==============================================================================
Convierte documentos Markdown (.md) a documentos Word (.docx) con diseño
profesional:
- Tipografía moderna (Segoe UI / Calibri) con jerarquía visual clara.
- Tablas estilizadas con encabezados en azul marino, filas alternadas y bordes sutiles.
- Bloques de código y diagramas ASCII en cajas sombreadas (Consolas) con borde lateral.
- Soporte para listas, negritas, cursivas, citas y símbolos matemáticos comunes.

Uso en consola:
  python3 scripts/md_to_docx.py --input docs/ESTANDAR_COMUNICACION_PC_ESP32.md --output-dir .
  python3 scripts/md_to_docx.py --input docs/ --output-dir .
==============================================================================
"""

import os
import sys
import re
import argparse
from pathlib import Path

try:
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn, nsdecls
except ImportError:
    print("Error: Se requiere la biblioteca python-docx.")
    print("Instalala ejecutando: pip install python-docx")
    sys.exit(1)


# =============================================================================
# PALETA DE COLORES CORPORATIVA / ACADÉMICA
# =============================================================================
COLOR_PRIMARY_DARK = RGBColor(15, 23, 42)      # Azul medianoche oscuro (#0F172A)
COLOR_PRIMARY_BLUE = RGBColor(3, 105, 161)     # Azul acero institucional (#0369A1)
COLOR_SECONDARY = RGBColor(71, 85, 105)        # Gris pizarra (#475569)
COLOR_BODY = RGBColor(30, 41, 59)              # Texto cuerpo legible (#1E293B)
COLOR_CODE = RGBColor(15, 23, 42)              # Código Consolas (#0F172A)
HEX_TABLE_HEADER = "1E293B"                    # Fondo encabezado tabla
HEX_TABLE_ALT_ROW = "F8FAFC"                   # Fondo fila alternada
HEX_BORDER = "CBD5E1"                          # Bordes de tabla (#CBD5E1)
HEX_CODE_BG = "F1F5F9"                         # Fondo bloque código (#F1F5F9)
HEX_CODE_LEFT_BORDER = "0284C7"                # Borde lateral código (#0284C7)


# =============================================================================
# UTILIDADES XML PARA DOCX (BORDES, SOMBREADOS Y MÁRGENES)
# =============================================================================
def set_cell_background(cell, hex_color):
    """Aplica color de fondo a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Establece márgenes internos (padding) de una celda en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_styling(table, col_widths=None):
    """Aplica bordes sutiles y alineación a la tabla."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Ancho automático proporcionado si se definió
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)


# =============================================================================
# LIMPIEZA Y FORMATEO DE TEXTO
# =============================================================================
def clean_markdown_math_symbols(text):
    """Reemplaza fórmulas TeX elementales por caracteres legibles en Word."""
    text = text.replace("$\\to$", "→").replace("$\\leftrightarrow$", "↔")
    text = text.replace("$\\le$", "≤").replace("$\\ge$", "≥")
    text = text.replace("$60^\\circ$", "60°").replace("$90^\\circ$", "90°").replace("$120^\\circ$", "120°")
    text = text.replace("$135^\\circ$", "135°").replace("$140^\\circ$", "140°").replace("$45^\\circ$", "45°")
    text = text.replace("$q_1$", "q₁").replace("$q_2$", "q₂").replace("$q$", "q")
    text = text.replace("$\\approx", "≈").replace("$$", "")
    return text

def add_formatted_runs(paragraph, text, base_font_name="Segoe UI", base_font_size=10.5, is_italic=False, color=COLOR_BODY):
    """
    Parsea fragmentos en negrita (**texto**), cursiva (*texto*) y código inline (`texto`)
    y los agrega como Runs con formato al párrafo.
    """
    text = clean_markdown_math_symbols(text)
    
    # Expresión regular que separa negritas, código inline y texto normal
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = base_font_name
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = color
        run.italic = is_italic

        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(194, 65, 12)  # Ámbar oscuro/naranja código (#C2410C)
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
        else:
            run.text = token


# =============================================================================
# CLASE PRINCIPAL DEL GENERADOR DOCX
# =============================================================================
class MarkdownToWordConverter:
    """Convierte un documento Markdown en un documento Word profesional."""

    def __init__(self, md_path, output_dir=None):
        self.md_path = Path(md_path)
        self.output_dir = Path(output_dir) if output_dir else self.md_path.parent
        self.output_docx = self.output_dir / f"{self.md_path.stem}.docx"
        self.doc = Document()
        self._configure_document_styles()

    def _configure_document_styles(self):
        """Configura márgenes estándar de 2.5 cm y tipografía base."""
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Configurar estilo Normal
        normal_style = self.doc.styles['Normal']
        normal_style.font.name = 'Segoe UI'
        normal_style.font.size = Pt(10.5)
        normal_style.font.color.rgb = COLOR_BODY
        normal_style.paragraph_format.line_spacing = 1.15
        normal_style.paragraph_format.space_after = Pt(4)

    def convert(self):
        """Lee el Markdown y genera el documento Word."""
        with open(self.md_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()

        i = 0
        n = len(lines)
        in_code_block = False
        code_lines = []

        while i < n:
            line = lines[i].rstrip('\r\n')
            stripped = line.strip()

            # -------------------------------------------------------------
            # 1. BLOQUES DE CÓDIGO Y DIAGRAMAS ASCII (```...```)
            # -------------------------------------------------------------
            if stripped.startswith("```"):
                if in_code_block:
                    # Cerrar bloque de código
                    self._render_code_box(code_lines)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lines = []
                i += 1
                continue

            if in_code_block:
                code_lines.append(line)
                i += 1
                continue

            # -------------------------------------------------------------
            # 2. LÍNEA HORIZONTAL DIVISORIA (---)
            # -------------------------------------------------------------
            if stripped in ("---", "***", "___"):
                self._render_divider()
                i += 1
                continue

            # -------------------------------------------------------------
            # 3. TABLAS DE MARKDOWN (| col1 | col2 |)
            # -------------------------------------------------------------
            if stripped.startswith("|") and stripped.endswith("|") and i + 1 < n and "---" in lines[i+1]:
                table_lines = []
                while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                self._render_markdown_table(table_lines)
                continue

            # -------------------------------------------------------------
            # 4. ENCABEZADOS (#, ##, ###, ####)
            # -------------------------------------------------------------
            if stripped.startswith("#"):
                self._render_heading(stripped)
                i += 1
                continue

            # -------------------------------------------------------------
            # 5. LISTAS CON VIÑETAS O NUMERADAS (*, -, 1.)
            # -------------------------------------------------------------
            if stripped.startswith(("* ", "- ")):
                content = stripped[2:].strip()
                p = self.doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                add_formatted_runs(p, content)
                i += 1
                continue

            num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
            if num_match:
                content = num_match.group(2)
                p = self.doc.add_paragraph(style='List Number')
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                add_formatted_runs(p, content)
                i += 1
                continue

            # -------------------------------------------------------------
            # 6. CITAS O NOTAS (> ...)
            # -------------------------------------------------------------
            if stripped.startswith(">"):
                quote_text = stripped[1:].strip()
                self._render_callout(quote_text)
                i += 1
                continue

            # -------------------------------------------------------------
            # 7. LÍNEA EN BLANCO
            # -------------------------------------------------------------
            if not stripped:
                i += 1
                continue

            # -------------------------------------------------------------
            # 8. PÁRRAFO DE TEXTO NORMAL
            # -------------------------------------------------------------
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, stripped)
            i += 1

        # Asegurar creación de directorio de salida
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(self.output_docx))
        return self.output_docx

    # -------------------------------------------------------------------------
    # RENDERIZADORES DE ELEMENTOS VISUALES
    # -------------------------------------------------------------------------
    def _render_heading(self, line):
        """Renderiza títulos con colores y jerarquía limpia."""
        if line.startswith("# "):
            text = line[2:].strip()
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_markdown_math_symbols(text))
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY_DARK

        elif line.startswith("## "):
            text = line[3:].strip()
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_markdown_math_symbols(text))
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY_BLUE

        elif line.startswith("### "):
            text = line[4:].strip()
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_markdown_math_symbols(text))
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = COLOR_BODY

        elif line.startswith("#### "):
            text = line[5:].strip()
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_markdown_math_symbols(text))
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = COLOR_SECONDARY

    def _render_divider(self):
        """Dibuja una barra divisoria elegante."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        p_border = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="{HEX_BORDER}"/>'
            f'</w:pBdr>'
        )
        p._p.get_or_add_pPr().append(p_border)

    def _render_callout(self, text):
        """Renderiza una caja de cita o nota destacada con borde lateral azul."""
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

        # Borde lateral izquierdo
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>'
            f'<w:top w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, text, is_italic=True, color=COLOR_SECONDARY)

        # Espacio después del callout
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    def _render_code_box(self, code_lines):
        """
        Renderiza un bloque de código o diagrama ASCII en una caja sombreada
        con fuente Consolas de 8.5 pt para que no se desalineen las columnas.
        """
        if not code_lines:
            return

        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.columns[0].width = Inches(6.5)
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

        # Borde sutil alrededor y acento en la izquierda
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_CODE_LEFT_BORDER}"/>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

        # Si el bloque contiene diagramas anchos, reducimos ligeramente la fuente para evitar quiebres
        max_line_len = max(len(l) for l in code_lines) if code_lines else 0
        font_size = 7.5 if max_line_len > 70 else 8.5

        for idx, line in enumerate(code_lines):
            run = p.add_run(line + ("\n" if idx < len(code_lines) - 1 else ""))
            run.font.name = "Consolas"
            run.font.size = Pt(font_size)
            run.font.color.rgb = COLOR_CODE

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def _render_markdown_table(self, table_lines):
        """Parsea una tabla en formato Markdown y la dibuja con estilos limpios."""
        if len(table_lines) < 2:
            return

        # Parsear filas
        raw_rows = []
        for line in table_lines:
            cells = [c.strip() for c in line.strip('|').split('|')]
            raw_rows.append(cells)

        header_row = raw_rows[0]
        # La segunda fila suele ser de alineación |:---|:---:|:---|
        data_rows = raw_rows[2:] if len(raw_rows) > 2 and "---" in raw_rows[1][0] else raw_rows[1:]

        num_cols = len(header_row)
        num_rows = 1 + len(data_rows)

        tbl = self.doc.add_table(rows=num_rows, cols=num_cols)
        set_table_styling(tbl)

        # 1. Fila de Encabezado (Azul medianoche y texto blanco en negrita)
        hdr_cells = tbl.rows[0].cells
        for col_idx, col_text in enumerate(header_row):
            if col_idx < len(hdr_cells):
                cell = hdr_cells[col_idx]
                set_cell_background(cell, HEX_TABLE_HEADER)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(clean_markdown_math_symbols(col_text))
                run.bold = True
                run.font.name = "Segoe UI Semibold"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

        # 2. Filas de Datos (Alternadas con sombreado muy suave)
        for row_idx, row_data in enumerate(data_rows):
            row_cells = tbl.rows[row_idx + 1].cells
            bg_color = HEX_TABLE_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_text in enumerate(row_data):
                if col_idx < len(row_cells):
                    cell = row_cells[col_idx]
                    set_cell_background(cell, bg_color)
                    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.1
                    add_formatted_runs(p, cell_text, base_font_size=9.5)

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)


# =============================================================================
# CLI PRINCIPAL
# =============================================================================
def main():
    parser = argparse.ArgumentParser(
        description="Convierte documentos Markdown (.md) a archivos Word (.docx) elegantes."
    )
    parser.add_argument(
        "--input", "-i",
        required=True,
        help="Ruta del archivo Markdown (.md) o carpeta con archivos .md a convertir."
    )
    parser.add_argument(
        "--output-dir", "-o",
        default=None,
        help="Carpeta destino donde guardar los archivos .docx (por defecto la misma o el root)."
    )

    args = parser.parse_args()
    input_path = Path(args.input)
    out_dir = Path(args.output_dir) if args.output_dir else None

    if not input_path.exists():
        print(f"Error: La ruta especificada no existe: {input_path}")
        sys.exit(1)

    files_to_convert = []
    if input_path.is_file() and input_path.suffix.lower() == ".md":
        files_to_convert.append(input_path)
    elif input_path.is_dir():
        files_to_convert.extend(list(input_path.glob("*.md")))

    if not files_to_convert:
        print(f"No se encontraron archivos .md para convertir en: {input_path}")
        sys.exit(0)

    print(f"Procesando {len(files_to_convert)} archivo(s) Markdown a Word...")
    for md_file in sorted(files_to_convert):
        converter = MarkdownToWordConverter(md_file, output_dir=out_dir)
        output_file = converter.convert()
        print(f"  [OK] Generado: {output_file}")

    print("\nConversion completada exitosamente.")


if __name__ == "__main__":
    main()
